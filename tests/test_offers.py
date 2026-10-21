import asyncio
import io
import json
import zipfile
from email.message import EmailMessage
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter

from app import db, offers, offers_api, pipeline
from app.docs import DocError, Document, read_document
from app.main import app
from app.models import Fact, Supplier
from app.offers import (auto_target, build_offer_data, extract_offer, match_supplier, numbers_in,
                        supplier_from_offer)
from app.pricing import compare, history, norm_unit, similarity

# Все документы ниже синтетические, нужны только тестам.
TEXT = """КП № 5 от 01.09.2026
ООО «Тест-Поставка» ИНН 7701234567
Тел.: +7 (812) 123-45-67, sales@test-postavka.ru, сайт test-postavka.ru
Цены с НДС. Минимальный заказ 100 кг.
Мука пшеничная в/с | мешок 25 кг | кг | 48,00
Сахар-песок | мешок 50 кг | кг | 1 061,50
"""


def raw(items=None, **over):
    base = {"is_offer": True, "supplier": {"name": "ООО «Тест-Поставка»", "inn": "7701234567",
            "website": "test-postavka.ru", "emails": ["sales@test-postavka.ru"],
            "phones": ["+7 (812) 123-45-67"]}, "offer_date": "2026-09-01", "currency": "RUB", "vat_included": True,
            "items": items if items is not None else [
                {"name": "Мука пшеничная в/с", "unit": "кг", "price": 48.0, "quote": "Мука пшеничная в/с | мешок 25 кг | кг | 48,00"},
                {"name": "Сахар-песок", "unit": "кг", "price": 1061.5, "quote": "Сахар-песок | мешок 50 кг | кг | 1 061,50"}]}
    base.update(over)
    return base


# --- чтение документов -----------------------------------------------------------

def test_read_text_docx_xlsx_eml():
    assert "Мука" in read_document("a.txt", TEXT.encode()).text

    xml = ('<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
           "<w:p><w:r><w:t>Мука пшеничная</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>48,00</w:t></w:r></w:p></w:body></w:document>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    d = read_document("kp.docx", buf.getvalue())
    assert "Мука пшеничная" in d.text and "48,00" in d.text and d.kind == "docx"

    wb = Workbook(); ws = wb.active; ws.append(["Мука", "кг", 48]); ws.append(["Сахар", "кг", 61.5])
    out = io.BytesIO(); wb.save(out)
    x = read_document("price.xlsx", out.getvalue())
    assert "Мука | кг | 48" in x.text and "61.5" in x.text

    msg = EmailMessage(); msg["From"] = "Иван <i@test-postavka.ru>"; msg["Subject"] = "КП"
    msg.set_content("Мука 48 руб/кг")
    e = read_document("m.eml", msg.as_bytes())
    assert "i@test-postavka.ru" in e.text and "Мука 48" in e.text


def test_encodings_are_not_garbled():
    body = "Мука пшеничная 48 руб/кг"
    # письмо без заголовка кодировки (как пишут некоторые почтовики) и текст в windows-1251
    raw_eml = ("From: a@b.ru\nSubject: КП\n\n" + body).encode("utf-8")
    assert body in read_document("m.eml", raw_eml).text
    assert body in read_document("m.txt", body.encode("cp1251")).text
    html_eml = ("From: a@b.ru\nContent-Type: text/html; charset=windows-1251\n\n<p>" + body + "</p>").encode("cp1251")
    assert body in read_document("h.eml", html_eml).text


def test_scanned_pdf_and_image_go_to_model_as_is():
    w = PdfWriter(); w.add_blank_page(100, 100); buf = io.BytesIO(); w.write(buf)
    d = read_document("scan.pdf", buf.getvalue())
    assert d.media and d.media["type"] == "document" and d.text == ""
    img = read_document("p.png", b"\x89PNG\r\n\x1a\n" + b"0" * 20)
    assert img.media["type"] == "image"


def test_bad_files_are_rejected():
    for name, content in [("a.exe", b"MZ"), ("a.txt", b""), ("a.docx", b"not a zip"), ("a.txt", b"x" * (9 * 1024 * 1024))]:
        with pytest.raises(DocError):
            read_document(name, content)


def test_zip_bomb_is_rejected(monkeypatch):
    monkeypatch.setattr("app.docs.MAX_UNPACKED", 1000)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "a" * 5000)
    with pytest.raises(DocError):
        read_document("bomb.docx", buf.getvalue())


# --- проверка по тексту ----------------------------------------------------------

def test_numbers_with_thousand_separators():
    n = numbers_in("цена 1 061,50 руб, 25 кг")
    assert 1061.5 in n and 25 in n and 61.5 not in n or 1061.5 in n


def test_invented_prices_are_dropped_real_ones_kept():
    items = [
        {"name": "Мука пшеничная в/с", "unit": "кг", "price": 48.0, "quote": "Мука пшеничная в/с | мешок 25 кг | кг | 48,00"},
        {"name": "Сахар-песок", "unit": "кг", "price": 1061.5, "quote": "Сахар-песок | мешок 50 кг | кг | 1 061,50"},
        {"name": "Соль", "unit": "кг", "price": 14.0, "quote": "Соль | мешок | кг | 14,00"},                     # выдумана
        {"name": "Мука пшеничная в/с", "unit": "кг", "price": 99.0, "quote": "Мука пшеничная в/с | мешок 25 кг | кг | 48,00"},  # цена не из цитаты
    ]
    data = build_offer_data(raw(items), Document(TEXT, None, "text"))
    assert [i["name"] for i in data["items"]] == ["Мука пшеничная в/с", "Сахар-песок"]
    assert data["rejected_items"] == 2 and data["checked_by_text"]
    assert all(i["verified"] for i in data["items"])


def test_table_rows_with_split_cells_are_accepted_by_fallback():
    doc = Document(TEXT.replace(" | ", "\t"), None, "text")            # ячейки с табами: точная цитата не совпадёт
    items = [{"name": "Мука пшеничная в/с", "unit": "кг", "price": 48.0, "quote": "Мука пшеничная в/с кг 48,00"}]
    assert len(build_offer_data(raw(items), doc)["items"]) == 1


def test_supplier_fields_are_verified_against_text():
    r = raw()
    r["supplier"]["emails"] += ["fake@nowhere.ru"]
    r["supplier"]["phones"] += ["+7 (999) 000-00-00"]
    r["supplier"]["inn"] = "1234567890"
    data = build_offer_data(r, Document(TEXT, None, "text"))
    assert data["supplier"]["emails"] == ["sales@test-postavka.ru"]
    assert data["supplier"]["phones"] == ["+7 (812) 123-45-67"]
    assert data["supplier"]["inn"] == "7701234567"          # ИНН берётся из текста, а не из ответа модели


def test_scan_items_are_marked_unverified():
    scan = Document("", {"type": "image", "media_type": "image/png", "data": "x"}, "image")
    data = build_offer_data(raw(), scan)
    assert data["checked_by_text"] is False and not any(i["verified"] for i in data["items"])


def test_not_an_offer_returns_none_and_bad_prices_are_skipped():
    assert build_offer_data({"is_offer": False}, Document(TEXT, None, "text")) is None
    items = [{"name": "X", "price": "много", "quote": "x"}, {"name": "Y", "price": -5, "quote": "y"}]
    assert build_offer_data(raw(items), Document(TEXT, None, "text"))["items"] == []


def test_extract_offer_uses_forced_tool_and_verifies():
    class FakeClient:
        def __init__(self): self.messages = SimpleNamespace(create=self._create); self.kwargs = None
        async def _create(self, **kw):
            self.kwargs = kw
            return SimpleNamespace(content=[SimpleNamespace(type="tool_use", input=raw())])
    c = FakeClient()
    data = asyncio.run(extract_offer(Document(TEXT, None, "text"), client=c))
    assert c.kwargs["tool_choice"]["name"] == "save_offer" and "<document>" in c.kwargs["messages"][0]["content"]
    assert len(data["items"]) == 2
    scan = Document("", {"type": "document", "media_type": "application/pdf", "data": "AAA"}, "pdf")
    asyncio.run(extract_offer(scan, client=c))
    assert c.kwargs["messages"][0]["content"][0]["type"] == "document"


# --- привязка к поставщику -------------------------------------------------------

def F(v):
    return Fact(value=v, quote="", url="")


def row(domain, **kw):
    return {"domain": domain, "data": Supplier(domain=domain, url=f"https://{domain}/", **kw).model_dump()}


INFO = {"name": "ООО «Тест-Поставка»", "inn": "7701234567", "website": None,
        "emails": ["sales@test-postavka.ru"], "phones": ["+7 (812) 123-45-67"]}


def test_match_by_inn_domain_phone_and_name():
    rows = [row("other.ru"), row("test-postavka.ru"), row("x.ru", inn=F("7701234567"))]
    c = match_supplier(INFO, rows)
    assert c[0]["score"] == 100 and c[0]["domain"] == "x.ru" and "ИНН совпал" in c[0]["reasons"]
    assert any(m["domain"] == "test-postavka.ru" and m["score"] == 85 for m in c)
    assert match_supplier({**INFO, "inn": None, "emails": [], "phones": ["+7 (812) 123-45-67"]},
                          [row("p.ru", phones=[F("+7 (812) 123-45-67")])])[0]["reasons"] == ["телефон совпал"]
    by_name = match_supplier({"name": "ООО Тест Поставка", "emails": [], "phones": []}, [row("n.ru", name=F("ООО Тест-Поставка"))])
    assert by_name and by_name[0]["score"] == 60 and auto_target(by_name) is None     # по одному названию не привязываем


def test_name_match_ignores_legal_forms_and_suffixes():
    from app.offers import names_overlap
    assert names_overlap("ООО «Образец-Мука»", "ООО «Образец-Мука» (вымышленный пример)")
    assert names_overlap("ООО «ПТК Экор»", "Экор (ООО «ПТК ЭКОР», производство — ООО «ПТК ЭКОР-ФИШ»)")
    assert names_overlap("Молоко-Опт", 'ООО "Молоко-Опт"')
    assert not names_overlap("ООО «Ромашка»", "ООО «Василёк»")
    assert not names_overlap("ООО", "ООО «Ромашка»")                          # одни формы собственности: не повод
    assert not names_overlap("Сыр", "Сырная лавка")                           # слишком короткое общее слово
    # по одному названию кандидат есть, но автопривязки нет
    c = match_supplier({"name": "ООО «Образец-Мука»", "emails": [], "phones": []},
                       [row("obrazets.example", name=F("ООО «Образец-Мука» (вымышленный пример)"))])
    assert c and c[0]["score"] == 60 and auto_target(c) is None


def test_free_mail_does_not_match_and_ambiguity_blocks_auto_assign():
    free = {"name": None, "inn": None, "emails": ["boss@gmail.com"], "phones": []}
    assert match_supplier(free, [row("gmail.com")]) == []
    two = [{"domain": "a", "score": 100}, {"domain": "b", "score": 95}]
    assert auto_target(two) is None and auto_target([{"domain": "a", "score": 100}]) == "a"
    assert auto_target([]) is None and auto_target([{"domain": "a", "score": 60}]) is None


def test_supplier_from_offer_without_site_gets_stable_domain():
    s = supplier_from_offer({"name": "ООО Ромашка", "emails": ["a@gmail.com"], "phones": [], "inn": None, "website": None})
    assert s.domain == "kp-ooo-romaska" and s.url == "" and s.name.value == "ООО Ромашка"
    s2 = supplier_from_offer(INFO)
    assert s2.domain == "test-postavka.ru"


# --- сравнение цен ---------------------------------------------------------------

def items(*rows):
    return {"items": [{"name": n, "unit": u, "price": p} for n, u, p in rows], "currency": "RUB", "vat_included": True}


def test_units_and_similarity():
    assert norm_unit("руб./кг") == "кг" and norm_unit("за 1 кг") == "кг" and norm_unit("КГ.") == "кг"
    assert similarity("Сыр Моцарелла 45% 2 кг", "сыр моцарелла 45%, 2 кг") >= 0.9
    assert similarity("Сыр 2 кг", "Сыр 5 кг") == 0.0            # другой вес = другая позиция
    assert similarity("Мука пшеничная 1 сорт", "Мука пшеничная высший сорт") == 0.0


def test_pack_written_into_name_in_one_offer_still_matches():
    old = {"items": [{"name": "Мука пшеничная в/с", "unit": "кг", "price": 48.0, "pack": "мешок 25 кг"},
                     {"name": "Сыр 2 кг", "unit": "кг", "price": 500.0}], "currency": "RUB"}
    new = {"items": [{"name": "Мука пшеничная в/с, мешок 25 кг", "unit": "кг", "price": 48.0, "pack": "мешок 25 кг"},
                     {"name": "Сыр 5 кг", "unit": "кг", "price": 480.0}], "currency": "RUB"}
    r = compare(old, new)
    by = {row["name"]: row for row in r["rows"]}
    assert by["Мука пшеничная в/с, мешок 25 кг"]["trend"] == "same"       # та же позиция, фасовка дописана в название
    assert by["Сыр 5 кг"]["trend"] == "new" and [i["name"] for i in r["removed"]] == ["Сыр 2 кг"]   # другой вес: другая позиция
    h = history([{"id": 1, "date": "2026-09-01", "data": old}, {"id": 2, "date": "2026-10-01", "data": new}])
    assert [len(c["series"]) for c in h if c["name"].startswith("Мука")] == [2]


def test_pack_phrases_are_stripped_but_meaningful_numbers_stay():
    s = similarity
    assert s("Моцарелла 45%, брус 2,3 кг", "Моцарелла 45%") >= 0.95
    assert s("Масло сливочное 82,5% (25 кг короб)", "Масло сливочное 82,5%") >= 0.95
    assert s("Мука в/с мешок 25 кг", "Мука в/с") >= 0.95
    assert s("Моцарелла 45%", "Моцарелла 20%") == 0.0              # жирность различает позиции
    assert s("Мука пшеничная 1 сорт", "Мука пшеничная высший сорт") == 0.0
    assert s("Сыр 2 кг", "Сыр 5 кг") == 0.0


def test_compare_detects_up_down_same_new_removed_unit_change():
    old = items(("Мука в/с", "кг", 48.0), ("Сахар", "кг", 61.0), ("Соль", "кг", 14.0), ("Масло", "л", 118.0), ("Дрожжи", "уп", 100.0))
    new = items(("Мука в/с", "кг", 52.0), ("Сахар", "кг", 58.0), ("Масло", "л", 118.0), ("Крупа", "кг", 40.0), ("Дрожжи", "кг", 200.0))
    r = compare(old, new)
    by = {row["name"]: row for row in r["rows"]}
    assert by["Мука в/с"]["trend"] == "up" and by["Мука в/с"]["delta_pct"] == 8.3
    assert by["Сахар"]["trend"] == "down" and by["Сахар"]["delta"] == -3.0
    assert by["Масло"]["trend"] == "same" and by["Крупа"]["trend"] == "new"
    assert by["Дрожжи"]["trend"] == "unit_changed" and by["Дрожжи"]["delta_pct"] is None
    assert [i["name"] for i in r["removed"]] == ["Соль"]
    assert r["summary"]["up"] == 1 and r["summary"]["down"] == 1 and r["summary"]["new"] == 1


def test_compare_without_previous_and_warnings():
    assert compare(None, items(("X", "кг", 1.0)))["has_previous"] is False
    old = {**items(("X", "кг", 10.0)), "vat_included": False}
    new = {**items(("X", "кг", 10.0)), "currency": "USD"}
    w = compare(old, new)["warnings"]
    assert len(w) == 2


def test_history_builds_series_per_item():
    offs = [{"id": 1, "date": "2026-09-01", "data": items(("Мука", "кг", 48.0), ("Сахар", "кг", 61.0))},
            {"id": 2, "date": "2026-10-01", "data": items(("мука", "кг", 52.0))}]
    h = {c["name"]: c for c in history(offs)}
    assert [p["price"] for p in h["Мука"]["series"]] == [48.0, 52.0] and h["Мука"]["change_pct"] == 8.3
    assert h["Сахар"]["change_pct"] is None


# --- API целиком -----------------------------------------------------------------

@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "k"); monkeypatch.setenv("TAVILY_API_KEY", "k")
    return TestClient(app)


def fake_extract(payload):
    async def _f(doc, client=None):
        return build_offer_data(payload, doc)
    return _f


def test_upload_matches_existing_supplier_and_tracks_prices(client, monkeypatch):
    db.save_supplier("test-postavka.ru", Supplier(domain="test-postavka.ru", url="https://test-postavka.ru/",
                                                          inn=F("7701234567")).model_dump())
    monkeypatch.setattr(offers_api, "extract_offer", fake_extract(raw()))
    r1 = client.post("/api/offers", files={"file": ("kp1.txt", TEXT.encode(), "text/plain")})
    assert r1.status_code == 200
    j1 = r1.json()
    assert j1["match"]["assigned"] == "test-postavka.ru" and j1["offer"]["diff"]["has_previous"] is False
    assert db.get_supplier_row("test-postavka.ru")["status"] == "quote"      # статус обновился сам

    text2 = TEXT.replace("48,00", "52,00").replace("01.09.2026", "01.10.2026")
    items2 = [{"name": "Мука пшеничная в/с", "unit": "кг", "price": 52.0, "quote": "Мука пшеничная в/с | мешок 25 кг | кг | 52,00"}]
    monkeypatch.setattr(offers_api, "extract_offer", fake_extract(raw(items2, offer_date="2026-10-01")))
    j2 = client.post("/api/offers", files={"file": ("kp2.txt", text2.encode(), "text/plain")}).json()
    row = j2["offer"]["diff"]["rows"][0]
    assert row["trend"] == "up" and row["prev_price"] == 48.0 and row["delta_pct"] == 8.3

    hist = client.get("/api/supplier/test-postavka.ru/offers").json()
    assert [o["filename"] for o in hist["offers"]] == ["kp2.txt", "kp1.txt"]      # новые сверху
    assert hist["history"][0]["change_pct"] == 8.3
    assert client.get(f"/api/offers/{j1['offer']['id']}/file").content == TEXT.encode()
    ov = client.get("/api/offers/overview").json()
    assert ov["items"][0]["count"] == 2 and ov["unassigned"] == 0


def test_unknown_supplier_stays_unassigned_until_user_decides(client, monkeypatch):
    monkeypatch.setattr(offers_api, "extract_offer", fake_extract(raw()))
    j = client.post("/api/offers", data={"text": TEXT}).json()
    assert j["match"]["assigned"] is None and j["offer"]["domain"] is None
    assert len(client.get("/api/offers/unassigned").json()["items"]) == 1
    assert client.post(f"/api/offers/{j['offer']['id']}/assign", json={}).status_code == 400
    done = client.post(f"/api/offers/{j['offer']['id']}/assign", json={"create": True}).json()
    assert done["offer"]["domain"] == "test-postavka.ru"
    assert client.get("/api/offers/unassigned").json()["items"] == []
    assert db.get_supplier_row("test-postavka.ru")["data"]["inn"]["value"] == "7701234567"
    assert client.post(f"/api/offers/{j['offer']['id']}/assign", json={"domain": "nope.ru"}).status_code == 400


def test_upload_errors_and_limits(client, monkeypatch):
    assert client.post("/api/offers").status_code == 400
    assert client.post("/api/offers", files={"file": ("a.exe", b"MZ", "application/x")}).status_code == 400
    monkeypatch.setattr(offers_api, "extract_offer", fake_extract({"is_offer": False}))
    assert client.post("/api/offers", data={"text": "привет"}).status_code == 422

    monkeypatch.setattr(offers_api, "extract_offer", fake_extract(raw()))
    monkeypatch.setenv("OFFER_DAILY_LIMIT", "2")          # одна попытка (422 выше) уже потратила ИИ и учтена
    assert client.post("/api/offers", data={"text": TEXT}).status_code == 200
    assert client.post("/api/offers", data={"text": TEXT}).status_code == 429

    async def boom(doc, client=None):
        raise RuntimeError("api down")
    monkeypatch.setenv("OFFER_DAILY_LIMIT", "50")
    monkeypatch.setattr(offers_api, "extract_offer", boom)
    assert client.post("/api/offers", data={"text": TEXT}).status_code == 502


def test_upload_without_keys_is_503_and_delete_removes_file(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", ""); monkeypatch.setenv("TAVILY_API_KEY", "")
    c = TestClient(app)
    assert c.post("/api/offers", data={"text": TEXT}).status_code == 503

    monkeypatch.setenv("ANTHROPIC_API_KEY", "k"); monkeypatch.setenv("TAVILY_API_KEY", "k")
    monkeypatch.setattr(offers_api, "extract_offer", fake_extract(raw()))
    j = c.post("/api/offers", files={"file": ("../../evil.txt", TEXT.encode(), "text/plain")}).json()
    assert j["offer"]["filename"] == "evil.txt"                                   # путь из имени вырезан
    assert c.delete(f"/api/offers/{j['offer']['id']}").status_code == 200
    assert c.get(f"/api/offers/{j['offer']['id']}/file").status_code == 404
    assert not list(offers_api._upload_dir().glob("*"))


def test_demo_offers_are_idempotent_and_show_price_changes():
    c = TestClient(app)          # без ключей: демо должно работать
    a = c.post("/api/offers/demo").json()
    b = c.post("/api/offers/demo").json()
    assert len(a["offers"]) == 2 and len(b["offers"]) == 2 and a["domain"] == "obrazets-muka.example"
    newest = a["offers"][-1]["diff"]                       # views идут от старых к новым
    by = {r["name"]: r for r in newest["rows"]}
    assert by["Мука пшеничная хлебопекарная высший сорт"]["trend"] == "up"
    assert by["Сахар-песок"]["trend"] == "down" and by["Мука пшеничная цельнозерновая"]["trend"] == "new"
    assert [r["name"] for r in newest["removed"]] == ["Соль пищевая выварочная"]
    assert all(i["verified"] for i in a["offers"][0]["data"]["items"])
    assert db.get_supplier_row("obrazets-muka.example")["data"]["url"] == ""
    assert "вымышленный" in db.get_supplier_row("obrazets-muka.example")["data"]["name"]["value"]
