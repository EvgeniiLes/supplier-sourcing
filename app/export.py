import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .pricing import norm_unit

STATUS_RU = {"new": "новый", "contacted": "написал", "quote": "получил КП",
             "agreed": "договорились", "rejected": "отказ"}

HEADERS = ["Название", "Сайт", "Оценка", "Статус", "Телефоны", "Email", "Адрес", "Регион работы",
           "Мин. заказ", "Цены на искомый товар", "Общие условия по ценам", "Документы", "Доставка",
           "Плюсы", "Минусы", "Источники", "Заметка"]


def _v(f) -> str:
    return f["value"] if f else ""


def _prices(it: dict) -> str:
    """Цены на искомый товар: из выдачи поиска, а для шорт-листа и старых данных — из карточки поставщика."""
    prices = it.get("query_prices")
    if prices is None:
        prices = it["data"].get("query_prices")
    if not prices:
        return "не найдено на сайте: запросить КП" if prices is not None else "не указано"
    def one(p: dict) -> str:
        unit = norm_unit(p.get("unit"))              # в старых снимках единица могла быть записана как «₽/кг»
        text = f"{p['name']}: {p['price']:g} ₽" + (f"/{unit}" if unit else "")
        if p.get("base_price") is not None:
            text += f" (≈{p['base_price']:g} ₽/{p['base_unit']}{', расчёт' if p.get('derived') else ''})"
        return text
    return "; ".join(one(p) for p in prices)


def rows(items: list[dict]) -> list[list]:
    out = []
    for it in items:
        d = it["data"]
        sc = it.get("score") or {}
        regions = [r["value"] for r in d.get("regions", [])]
        if d.get("nationwide"):
            regions.append("вся Россия")
        out.append([
            _v(d.get("name")) or d["domain"], d["url"], sc.get("score", ""),
            STATUS_RU.get(it.get("status", "new"), ""),
            "; ".join(p["value"] for p in d.get("phones", [])),
            "; ".join(e["value"] for e in d.get("emails", [])),
            _v(d.get("address")), "; ".join(regions),
            _v(d.get("min_order")) or "не указано", _prices(it), _v(d.get("price_info")) or "не указано",
            "; ".join(c["value"] for c in d.get("certificates", [])) or "не указано",
            _v(d.get("delivery")) or "не указано",
            "; ".join(sc.get("pros", [])), "; ".join(sc.get("cons", [])),
            "\n".join(d.get("sources", [])[:5]), it.get("note", ""),
        ])
    return out


def to_csv(items: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(HEADERS)
    w.writerows(rows(items))
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM, чтобы Excel открыл кириллицу


def to_xlsx(items: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Поставщики"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D50")
    for r in rows(items):
        ws.append(r)
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = 14 if h in ("Оценка", "Статус") else 32
    ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
